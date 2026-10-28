"""Geração da aba mensal "ALTERAÇÕES CNES" a partir do modelo (planilha_modelo.xlsx)."""
from __future__ import annotations

import io
import shutil
import subprocess
from copy import copy, deepcopy
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .paths import resource_path
from .validators import MESES

COLS = range(1, 9)            # A..H
HEADER_LAST_ROW = 6
FIRST_DATA_ROW = 7
MIN_DATA_ROWS = 18            # linhas em branco mínimas, como no modelo
TPL_FOOTER_FIRST = 25         # 1ª linha do rodapé no modelo
TPL_FOOTER_LAST = 33
TPL_SHEET = "MODELO"


def sheet_name(ano: int, mes: int) -> str:
    return f"{MESES[mes - 1].capitalize()}{ano}"


def _copy_style(src, dst) -> None:
    dst.font = copy(src.font)
    dst.border = copy(src.border)
    dst.fill = copy(src.fill)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)
    dst.alignment = copy(src.alignment)


def _load_template() -> tuple[openpyxl.Workbook, Worksheet]:
    wb = openpyxl.load_workbook(resource_path("assets", "planilha_modelo.xlsx"))
    return wb, wb[TPL_SHEET]


def build_month_sheet(
    ws: Worksheet,
    *,
    municipio: str,
    linhas: list[dict[str, Any]],
    realizado_por: str,
    observacoes: str = "",
) -> None:
    """Escreve o conteúdo de uma aba mensal em `ws` (que deve estar vazia).

    Cada item de `linhas`: cnes, estabelecimento, competencia (texto), data_envio (date|None),
    e UM dos campos: estab, inclusao, exclusao, outras (texto).
    """
    tpl_wb, tpl = _load_template()

    # larguras de coluna
    for _, dim in tpl.column_dimensions.items():
        for idx in range(dim.min, min(dim.max, 12) + 1):
            ws.column_dimensions[get_column_letter(idx)].width = dim.width

    # cabeçalho (linhas 1-6)
    for r in range(1, HEADER_LAST_ROW + 1):
        ws.row_dimensions[r].height = tpl.row_dimensions[r].height
        for c in COLS:
            s, d = tpl.cell(r, c), ws.cell(r, c)
            d.value = s.value
            _copy_style(s, d)
    for rng in tpl.merged_cells.ranges:
        if rng.max_row <= HEADER_LAST_ROW:
            ws.merge_cells(str(rng))
    ws["A3"] = f"MUNICÍPIO: {municipio.upper()}"

    # linhas de dados
    n = max(MIN_DATA_ROWS, len(linhas))
    tpl_first, tpl_mid, tpl_last = FIRST_DATA_ROW, FIRST_DATA_ROW + 1, TPL_FOOTER_FIRST - 1
    for i in range(n):
        r = FIRST_DATA_ROW + i
        src_row = tpl_first if i == 0 else (tpl_last if i == n - 1 else tpl_mid)
        ws.row_dimensions[r].height = 22.5
        for c in COLS:
            _copy_style(tpl.cell(src_row, c), ws.cell(r, c))
            ws.cell(r, c).number_format = "General"
        if i < len(linhas):
            it = linhas[i]
            ws.cell(r, 1, int(it["cnes"]) if str(it["cnes"]).isdigit() else it["cnes"])
            ws.cell(r, 2, it["estabelecimento"])
            ws.cell(r, 3, it["competencia"])
            ws.cell(r, 4, it.get("data_envio"))
            ws.cell(r, 4).number_format = "dd/mm/yyyy"
            for col, key in ((5, "estab"), (6, "inclusao"), (7, "exclusao"), (8, "outras")):
                if it.get(key):
                    ws.cell(r, col, it[key])

    # rodapé: copia as linhas 25..33 do modelo logo após os dados
    offset = (FIRST_DATA_ROW + n) - TPL_FOOTER_FIRST
    for r in range(TPL_FOOTER_FIRST, TPL_FOOTER_LAST + 1):
        nr = r + offset
        ws.row_dimensions[nr].height = tpl.row_dimensions[r].height
        for c in COLS:
            s, d = tpl.cell(r, c), ws.cell(nr, c)
            d.value = s.value
            _copy_style(s, d)
    for rng in tpl.merged_cells.ranges:
        if rng.min_row >= TPL_FOOTER_FIRST:
            ws.merge_cells(
                start_row=rng.min_row + offset, end_row=rng.max_row + offset,
                start_column=rng.min_col, end_column=rng.max_col,
            )
    ws.cell(27 + offset, 1, f"ALTERAÇÕES REALIZADAS POR: {realizado_por.upper()}")
    ws.cell(28 + offset, 1, "OBSERVAÇÕES:  " + observacoes)

    # configuração de página e logo
    ws.page_setup.orientation = tpl.page_setup.orientation
    ws.page_setup.paperSize = tpl.page_setup.paperSize
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    for attr in ("left", "right", "top", "bottom", "header", "footer"):
        setattr(ws.page_margins, attr, getattr(tpl.page_margins, attr))
    ws.print_area = f"A1:H{TPL_FOOTER_LAST + offset}"
    for img in tpl._images:
        new = XLImage(io.BytesIO(img._data()))
        new.anchor = deepcopy(img.anchor)
        ws.add_image(new)
    tpl_wb.close()


def montar_linhas(registros: Iterable[dict[str, Any]], estab: dict[str, Any],
                  competencia_texto: str) -> list[dict[str, Any]]:
    """Transforma registros do banco em linhas da planilha (ordem: estab., inclusão, exclusão, outras)."""
    ordem = {"estab": 0, "inclusao": 1, "exclusao": 2, "outras": 3}
    out: list[dict[str, Any]] = []
    for r in registros:
        item = {
            "cnes": estab["cnes"],
            "estabelecimento": estab["nome_fantasia"],
            "competencia": competencia_texto,
            "data_envio": _to_date(r.get("data_envio")),
        }
        if r["categoria"] == "ESTABELECIMENTO":
            item["estab"] = r.get("descricao_planilha") or "Alteração no cadastro do estabelecimento"
            item["_k"] = "estab"
        elif r["operacao"] == "INCLUSAO":
            item["inclusao"] = r["nome_profissional"]
            item["_k"] = "inclusao"
        elif r["operacao"] == "EXCLUSAO":
            item["exclusao"] = r["nome_profissional"]
            item["_k"] = "exclusao"
        else:
            desc = (r.get("descricao_planilha") or "").strip()
            item["outras"] = f"{r['nome_profissional']} ({desc})" if desc else r["nome_profissional"]
            item["_k"] = "outras"
        out.append(item)
    out.sort(key=lambda i: (ordem[i["_k"]], str(i.get("inclusao") or i.get("exclusao")
                                                 or i.get("outras") or i.get("estab")).upper()))
    for i in out:
        i.pop("_k")
    return out


def _to_date(v: Any) -> date | None:
    if isinstance(v, date):
        return v
    if isinstance(v, str) and v:
        try:
            return datetime.strptime(v[:10], "%Y-%m-%d").date()
        except ValueError:
            return None
    return None


def write_workbook(
    path: str | Path,
    *,
    ano: int,
    mes: int,
    municipio: str,
    linhas: list[dict[str, Any]],
    realizado_por: str,
    observacoes: str = "",
    anexar_existente: bool = False,
) -> tuple[Path, bool]:
    """Grava a aba do mês em `path` (.xlsx). Retorna (caminho, substituiu_aba_existente).

    Se `anexar_existente` e o arquivo existir, a aba do mês é adicionada (ou substituída)
    nesse arquivo, preservando as demais abas.
    """
    path = Path(path)
    name = sheet_name(ano, mes)
    replaced = False
    if anexar_existente and path.exists():
        wb = openpyxl.load_workbook(path)
        if name in wb.sheetnames:
            idx = wb.sheetnames.index(name)
            del wb[name]
            ws = wb.create_sheet(name, idx)
            replaced = True
        else:
            ws = wb.create_sheet(name, 0)  # mais recente primeiro, como na planilha original
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = name
    build_month_sheet(ws, municipio=municipio, linhas=linhas,
                      realizado_por=realizado_por, observacoes=observacoes)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path, replaced


def _libreoffice() -> str | None:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if exe:
        return exe
    for cand in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                 r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if Path(cand).exists():
            return cand
    return None


def libreoffice_disponivel() -> bool:
    return _libreoffice() is not None


def converter_para_xls(xlsx_path: str | Path) -> Path | None:
    """Converte .xlsx em .xls usando o LibreOffice, se instalado. Retorna None se indisponível."""
    exe = _libreoffice()
    if not exe:
        return None
    xlsx_path = Path(xlsx_path)
    subprocess.run(
        [exe, "--headless", "--convert-to", "xls", "--outdir", str(xlsx_path.parent), str(xlsx_path)],
        check=True, capture_output=True, timeout=180,
    )
    out = xlsx_path.with_suffix(".xls")
    return out if out.exists() else None
