"""Interface gráfica (customtkinter) da aplicação FlexCNES, com tema claro/escuro."""
from __future__ import annotations

import os
import queue
import re
import subprocess
import sys
import threading
import tkinter as tk
from datetime import date
from functools import lru_cache
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

import customtkinter as ctk

from . import __version__, cnes_import, planilha, settings
from .db import AuthError, Database
from .paths import app_db_path, cnes_db_path, default_output_dir, resource_path
from .security import password_problems
from .service import OPERACOES, Service, Session, competencia_key
from .settings import THEMES
from .validators import MESES, format_date

APP_TITLE = "FlexCNES"
APP_SUBTITLE = "Preenchimento das fichas e planilha CNES"

# ------------------------------------------------------------------------------
# Definição dos campos dos formulários
# ------------------------------------------------------------------------------
RACA_COR = ["01 - Branca", "02 - Preta", "03 - Parda", "04 - Amarela", "05 - Indígena",
            "99 - Sem informação"]

# (chave, rótulo, tipo, largura, opções)
F20_IDENT = [
    ("cadastramento", "Cadastramento", "radio", 0, [("SUS", "SUS"), ("NAO_SUS", "Não SUS")]),
    ("pis", "34.1.2 PIS/PASEP", "entry", 16, None),
    ("cpf", "34.1.3 CPF", "entry", 16, None),
    ("cns", "34.1.4 Número CNS", "entry", 18, None),
    ("sexo", "34.1.5 Sexo", "radio", 0, [("M", "M"), ("F", "F")]),
    ("mae", "34.1.6 Nome da mãe", "entry", 40, None),
    ("pai", "34.1.7 Nome do pai", "entry", 40, None),
    ("data_nasc", "34.1.8 Data de nascimento", "date", 12, None),
    ("mun_nasc", "34.1.9 Município de nascimento", "entry", 28, None),
    ("ibge_nasc", "34.1.10 Cód. IBGE do município", "entry", 10, None),
    ("uf_nasc", "34.1.11 UF", "entry", 4, None),
    ("raca_cor", "34.1.12 Raça/Cor", "combo", 16, RACA_COR),
    ("cert_tipo", "34.1.13 Certidão/Tipo", "entry", 6, None),
    ("cartorio", "34.1.14 Nome do cartório", "entry", 40, None),
    ("livro", "34.1.15 Livro", "entry", 10, None),
    ("fls", "34.1.16 Fls", "entry", 6, None),
    ("termo", "34.1.17 Termo", "entry", 10, None),
    ("cert_data", "34.1.18 Data de emissão (certidão)", "date", 12, None),
    ("rg_num", "34.1.19 Nº Identidade", "entry", 18, None),
    ("rg_uf", "34.1.20 UF", "entry", 4, None),
    ("rg_orgao", "34.1.21 Órgão emissor", "entry", 14, None),
    ("rg_data", "34.1.22 Data de emissão (identidade)", "date", 12, None),
    ("nacionalidade", "34.1.23 Nacionalidade", "radio", 0, [("B", "Brasileiro"), ("E", "Estrangeiro")]),
    ("pais_origem", "34.1.24 País de origem", "entry", 8, None),
    ("data_entrada", "34.1.25 Data de entrada (no país)", "date", 12, None),
    ("data_naturalizacao", "34.1.26 Data de naturalização", "date", 12, None),
    ("portaria", "34.1.27 Nº da portaria", "entry", 12, None),
    ("titulo", "34.1.28 Nº título de eleitor", "entry", 16, None),
    ("zona", "34.1.29 Zona", "entry", 6, None),
    ("secao", "34.1.30 Seção", "entry", 6, None),
    ("ctps_num", "34.1.31 CTPS número", "entry", 10, None),
    ("ctps_serie", "34.1.32 Série", "entry", 8, None),
    ("ctps_uf", "34.1.33 UF", "entry", 4, None),
    ("ctps_data", "34.1.34 Data de emissão (CTPS)", "date", 12, None),
    ("escolaridade", "34.1.35 Escolaridade (código)", "entry", 6, None),
    ("sit_conjugal", "34.1.36 Sit. familiar/conjugal (cód.)", "entry", 6, None),
    ("freq_escola", "34.1.37 Frequenta escola?", "radio", 0, [("S", "Sim"), ("N", "Não")]),
]
F20_RESID = [
    ("tipo_logradouro", "34.2.1 Tipo logradouro", "entry", 12, None),
    ("logradouro", "34.2.2 Logradouro", "entry", 36, None),
    ("numero", "34.2.3 Número", "entry", 8, None),
    ("complemento", "34.2.4 Complemento", "entry", 16, None),
    ("bairro", "34.2.5 Bairro/Distrito", "entry", 20, None),
    ("mun_res", "34.2.6 Município de residência", "entry", 28, None),
    ("ibge_res", "34.2.7 Cód. IBGE do município", "entry", 10, None),
    ("uf_res", "34.2.8 UF", "entry", 4, None),
    ("cep", "34.2.9 CEP", "entry", 12, None),
    ("telefone", "34.2.10 Telefone", "entry", 16, None),
]
F20_BANCO = [
    ("banco_cod", "34.3.1 Banco – código", "entry", 6, None),
    ("banco_nome", "Banco – nome", "entry", 22, None),
    ("ag_cod", "34.3.2 Agência – código", "entry", 8, None),
    ("ag_nome", "Agência – nome", "entry", 22, None),
    ("conta", "34.3.3 Conta corrente", "entry", 18, None),
]
F21_VINC = [
    ("registro", "34.4.1 Registro no conselho", "entry", 16, None),
    ("orgao", "34.4.2 Órgão emissor", "entry", 12, None),
    ("sus", "34.4.3 Atendimento ao SUS", "radio", 0, [("S", "Sim"), ("N", "Não")]),
    ("vinc_cod", "34.4.4 Vinculação – cód.", "entry", 6, None),
    ("vinc_desc", "Vinculação – descrição", "entry", 24, None),
    ("tipo_cod", "Tipo – cód.", "entry", 6, None),
    ("tipo_desc", "Tipo – descrição", "entry", 24, None),
    ("sub_cod", "Sub-tipo – cód.", "entry", 6, None),
    ("sub_desc", "Sub-tipo – descrição", "entry", 20, None),
    ("cbo_cod", "34.4.5 CBO – código", "entry", 10, None),
    ("cbo_desc", "CBO – descrição", "entry", 36, None),
    ("ch_amb", "34.4.6 CH ambulatorial", "entry", 6, None),
    ("ch_hosp", "CH hospitalar", "entry", 6, None),
    ("ch_outros", "CH outros", "entry", 6, None),
    ("data_entrada", "34.4.7 Data de entrada", "date", 12, None),
    ("data_deslig", "34.4.8 Data de desligamento", "date", 12, None),
    ("motivo_cod", "34.4.9 Motivo – cód.", "entry", 6, None),
    ("motivo_desc", "Motivo – descrição", "entry", 30, None),
]
F_ESTAB = [
    ("cnes", "CNES (7 dígitos)", "entry", 12, None),
    ("nome_fantasia", "Nome fantasia", "entry", 50, None),
    ("razao_social", "Razão social", "entry", 50, None),
    ("cnpj", "CNPJ", "entry", 20, None),
    ("municipio", "Município", "entry", 28, None),
    ("uf", "UF", "entry", 4, None),
    ("ibge", "Cód. IBGE do município", "entry", 12, None),
]


def open_path(path: Path) -> None:
    """Abre pasta/arquivo no gerenciador padrão do sistema."""
    try:
        if sys.platform.startswith("win"):
            os.startfile(str(path))  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.Popen(["open", str(path)])
        else:
            subprocess.Popen(["xdg-open", str(path)])
    except Exception:
        pass

# ------------------------------------------------------------------------------
# Tema, ícone e helpers de widgets
# ------------------------------------------------------------------------------
_TREE_COLORS = {
    "Light": dict(bg="#ffffff", fg="#1a1a1a", head="#dbdbdb", sel="#3b8ed0", selfg="#ffffff"),
    "Dark": dict(bg="#2b2b2b", fg="#dce4ee", head="#3a3a3a", sel="#1f6aa5", selfg="#ffffff"),
}


def tree_colors() -> dict:
    return _TREE_COLORS["Dark" if ctk.get_appearance_mode() == "Dark" else "Light"]


def style_native(master) -> None:
    """Ajusta widgets ttk (Treeview) ao tema atual; chamar de novo ao trocar o tema."""
    c = tree_colors()
    st = ttk.Style(master)
    try:
        st.theme_use("clam")
    except Exception:
        pass
    st.configure("Treeview", background=c["bg"], fieldbackground=c["bg"], foreground=c["fg"],
                 rowheight=26, borderwidth=0, relief="flat")
    st.configure("Treeview.Heading", background=c["head"], foreground=c["fg"], relief="flat",
                 borderwidth=0, padding=(6, 4))
    st.map("Treeview", background=[("selected", c["sel"])], foreground=[("selected", c["selfg"])])
    st.map("Treeview.Heading", background=[("active", c["head"])])


def set_icon(win) -> None:
    """Aplica assets/icon.ico (Windows) ou icon.png. O CTk redefine o ícone após ~200 ms."""
    ico, png = resource_path("assets", "icon.ico"), resource_path("assets", "icon.png")

    def apply():
        try:
            if sys.platform.startswith("win") and ico.exists():
                win.iconbitmap(str(ico))
            elif png.exists():
                win._icon_img = tk.PhotoImage(file=str(png))
                win.iconphoto(True, win._icon_img)
        except Exception:
            pass
    apply()
    try:
        win.after(250, apply)
    except Exception:
        pass


@lru_cache(maxsize=None)
def font(size: int = 13, bold: bool = False):
    return ctk.CTkFont(size=size, weight="bold" if bold else "normal")


def px(chars: int) -> int:
    """Converte largura em caracteres (usada nas listas de campos) para pixels."""
    return int(chars * 8.2) + 16


def button(master, text: str, command, kind: str = "primary", **kw):
    kw.setdefault("width", max(90, 28 + 8 * len(text)))
    if kind == "secondary":
        kw.update(fg_color=("gray78", "gray28"), hover_color=("gray68", "gray38"),
                  text_color=("gray10", "gray92"))
    elif kind == "danger":
        kw.update(fg_color=("#c0392b", "#a93226"), hover_color=("#992d22", "#8a281d"))
    return ctk.CTkButton(master, text=text, command=command, **kw)


def label(master, text: str = "", **kw):
    return ctk.CTkLabel(master, text=text, **kw)


def entry(master, var, chars: int = 20, **kw):
    return ctk.CTkEntry(master, textvariable=var, width=px(chars), **kw)


def combo(master, var, values, chars: int = 16, **kw):
    return ctk.CTkComboBox(master, variable=var, values=list(values), width=px(chars) + 14, **kw)


def radio(master, text: str, value, var, **kw):
    return ctk.CTkRadioButton(master, text=text, value=value, variable=var,
                              width=34 + 8 * len(text), **kw)


def check(master, text: str, var, **kw):
    return ctk.CTkCheckBox(master, text=text, variable=var, onvalue=True, offvalue=False,
                           width=34 + 8 * len(text), **kw)


def theme_menu(master, app: "App"):
    """Seletor Sistema / Claro / Escuro."""
    cur = next((k for k, v in THEMES.items() if v == app.theme), "Sistema")
    m = ctk.CTkOptionMenu(master, values=list(THEMES), width=110, command=app.set_theme)
    m.set(cur)
    return m


class TreeBox(ctk.CTkFrame):
    """ttk.Treeview (tabela) com barra de rolagem do CTk, estilizada pelo tema."""

    def __init__(self, master, columns, heads, height: int = 8):
        super().__init__(master, fg_color="transparent")
        self.tree = ttk.Treeview(self, columns=columns, show="headings", height=height,
                                 selectmode="browse")
        for c, t, w in heads:
            self.tree.heading(c, text=t)
            self.tree.column(c, width=w)
        sb = ctk.CTkScrollbar(self, command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")


def _auto_date(var: tk.StringVar) -> None:
    """Formata 8 dígitos digitados/colados como dd/mm/aaaa."""
    def cb(*_):
        t = var.get()
        d = re.sub(r"\D", "", t)
        if t == d and len(d) == 8:
            var.set(f"{d[:2]}/{d[2:4]}/{d[4:]}")
    var.trace_add("write", cb)


class FormPanel(ctk.CTkFrame):
    """Formulário em grade gerado a partir de uma lista de campos."""

    def __init__(self, master, fields, per_row: int = 3):
        super().__init__(master, fg_color="transparent")
        self.fields = fields
        self.vars: dict[str, tk.StringVar] = {}
        self.kinds: dict[str, str] = {}
        for i, (key, text, kind, width, opts) in enumerate(fields):
            r, c = divmod(i, per_row)
            cell = ctk.CTkFrame(self, fg_color="transparent")
            cell.grid(row=r, column=c, sticky="w", padx=6, pady=4)
            label(cell, text, anchor="w", font=font(12)).pack(anchor="w")
            var = tk.StringVar()
            self.vars[key] = var
            self.kinds[key] = kind
            if kind == "radio":
                row = ctk.CTkFrame(cell, fg_color="transparent")
                row.pack(anchor="w", pady=(2, 0))
                for value, t in opts:
                    radio(row, t, value, var).pack(side="left", padx=(0, 6))
            elif kind == "combo":
                combo(cell, var, opts, width).pack(anchor="w")
            else:
                entry(cell, var, width).pack(anchor="w")
                if kind == "date":
                    _auto_date(var)

    def get(self) -> dict[str, str]:
        out = {}
        for key, var in self.vars.items():
            v = var.get().strip()
            if self.kinds[key] == "combo" and " - " in v:
                v = v.split(" - ", 1)[0].strip()
            out[key] = v
        return out

    def set(self, data: dict) -> None:
        for key, var in self.vars.items():
            var.set(str(data.get(key, "") or ""))

    def clear(self) -> None:
        self.set({})


class ScrollFrame(ctk.CTkScrollableFrame):
    """Área rolável; `inner` é o próprio frame."""

    @property
    def inner(self):
        return self


class Modal(ctk.CTkToplevel):
    def __init__(self, master, title: str):
        super().__init__(master)
        self.title(title)
        self.transient(master.winfo_toplevel())
        self.resizable(False, False)
        self.result = None
        self.bind("<Escape>", lambda e: self.destroy())
        set_icon(self)

    def run(self):
        self.update_idletasks()
        try:
            self.wait_visibility()
            self.grab_set()
        except tk.TclError:
            pass
        self.wait_window(self)
        return self.result


class EstabDialog(Modal):
    def __init__(self, master, main: "MainFrame", est_id: int | None = None, first: bool = False):
        super().__init__(master, "Cadastro do estabelecimento")
        self.main, self.est_id = main, est_id
        body = ctk.CTkFrame(self, fg_color="transparent")
        body.pack(padx=16, pady=14)
        if first:
            label(body, "Cadastre o estabelecimento de saúde. Estes dados serão reutilizados em todas "
                        "as fichas (CNES e nome fantasia) e na planilha (município). Você poderá "
                        "editá-los depois.", wraplength=560, justify="left", anchor="w"
                  ).pack(anchor="w", pady=(0, 10))
        self.form = FormPanel(body, F_ESTAB, per_row=2)
        self.form.pack()
        if est_id:
            self.form.set(dict(main.app.db.get_estabelecimento(est_id)))
        btns = ctk.CTkFrame(body, fg_color="transparent")
        btns.pack(fill="x", pady=(12, 0))
        self.btn_busca = button(btns, "Buscar pelo CNES na base", self.buscar, kind="secondary")
        self.btn_busca.pack(side="left")
        button(btns, "Cancelar", self.destroy, kind="secondary").pack(side="right")
        button(btns, "Salvar", self.salvar).pack(side="right", padx=8)

    def buscar(self):
        base = self.main.app.cnes_base
        if not base.available():
            messagebox.showinfo("Base CNES", "A base CNES ainda não foi importada.\n"
                                "Use a aba “Base CNES” ou preencha manualmente.", parent=self)
            return
        e = base.estabelecimento(self.form.vars["cnes"].get())
        if not e:
            messagebox.showinfo("Base CNES", "CNES não encontrado na base importada.", parent=self)
            return
        self.form.set({"cnes": e["cnes"], "nome_fantasia": e.get("fantasia", ""),
                       "razao_social": e.get("razao", ""), "cnpj": e.get("cnpj", ""),
                       "municipio": e.get("municipio", ""), "uf": e.get("uf", ""),
                       "ibge": e.get("ibge", "")})

    def salvar(self):
        try:
            self.result = self.main.app.db.save_estabelecimento(self.form.get(), self.est_id)
        except ValueError as e:
            messagebox.showerror("Estabelecimento", str(e), parent=self)
            return
        self.destroy()


class CboSearch(Modal):
    def __init__(self, master, base: cnes_import.CnesBase):
        super().__init__(master, "Buscar CBO")
        self.base = base
        body = ctk.CTkFrame(self, fg_color="transparent")
        body.pack(padx=14, pady=12)
        self.q = tk.StringVar()
        e = entry(body, self.q, 56, placeholder_text="Digite ao menos 2 letras da ocupação")
        e.pack(anchor="w")
        e.bind("<KeyRelease>", lambda ev: self.refresh())
        c = tree_colors()
        self.lb = tk.Listbox(body, width=70, height=14, bg=c["bg"], fg=c["fg"], borderwidth=0,
                             highlightthickness=0, selectbackground=c["sel"],
                             selectforeground=c["selfg"], activestyle="none")
        self.lb.pack(pady=8, fill="x")
        self.lb.bind("<Double-Button-1>", lambda ev: self.pick())
        self.items: list[tuple[str, str]] = []
        button(body, "Usar selecionado", self.pick).pack(anchor="e")

    def refresh(self):
        t = self.q.get().strip()
        self.items = self.base.search_cbo(t) if len(t) >= 2 else []
        self.lb.delete(0, "end")
        for cod, desc in self.items:
            self.lb.insert("end", f"{cod} – {desc}")

    def pick(self):
        sel = self.lb.curselection()
        if sel:
            self.result = self.items[sel[0]]
            self.destroy()


# ------------------------------------------------------------------------------
# Telas de autenticação
# ------------------------------------------------------------------------------
class AuthFrame(ctk.CTkFrame):
    def __init__(self, master, app: "App", first: bool):
        super().__init__(master, fg_color="transparent")
        self.app, self.first = app, first
        theme_menu(self, app).place(relx=1.0, x=-14, y=14, anchor="ne")
        box = ctk.CTkFrame(self, corner_radius=14)
        box.place(relx=0.5, rely=0.45, anchor="center")
        inner = ctk.CTkFrame(box, fg_color="transparent")
        inner.pack(padx=36, pady=28)
        try:
            from PIL import Image
            _img = Image.open(resource_path("assets", "icon.png"))
            self.logo = ctk.CTkImage(light_image=_img, dark_image=_img, size=(80, 80))
            ctk.CTkLabel(inner, image=self.logo, text="").pack(pady=(0, 8))
        except Exception:
            pass  # sem ícone, o login continua funcionando
        label(inner, APP_TITLE, font=font(26, True)).pack()
        label(inner, APP_SUBTITLE, font=font(13)).pack(pady=(0, 10))
        label(inner, ("Primeiro acesso: crie o usuário administrador." if first
                      else "Entre com seu usuário e senha."), wraplength=380).pack(pady=(0, 14))
        self.nome, self.user = tk.StringVar(), tk.StringVar()
        self.pw, self.pw2 = tk.StringVar(), tk.StringVar()
        form = ctk.CTkFrame(inner, fg_color="transparent")
        form.pack()
        rows = ([("Nome completo", self.nome, False)] if first else []) + [
            ("Usuário", self.user, False), ("Senha", self.pw, True)] + (
            [("Confirmar senha", self.pw2, True)] if first else [])
        self.entries = []
        for i, (text, var, secret) in enumerate(rows):
            label(form, text).grid(row=i, column=0, sticky="e", padx=8, pady=5)
            e = ctk.CTkEntry(form, textvariable=var, width=270, show="•" if secret else "")
            e.grid(row=i, column=1, pady=5)
            e.bind("<Return>", lambda ev: self.submit())
            self.entries.append(e)
        button(inner, "Criar usuário e entrar" if first else "Entrar", self.submit,
               width=220, height=34).pack(pady=(18, 0))
        self.entries[0].focus_set()

    def submit(self):
        db = self.app.db
        try:
            if self.first:
                if self.pw.get() != self.pw2.get():
                    raise AuthError("As senhas não conferem.")
                probs = password_problems(self.pw.get())
                if probs:
                    raise AuthError(" ".join(probs))
                db.create_user(self.user.get(), self.nome.get(), self.pw.get())
            row = db.authenticate(self.user.get(), self.pw.get())
        except AuthError as e:
            messagebox.showerror("Acesso", str(e), parent=self)
            return
        self.app.on_authenticated(Session(row["id"], row["usuario"], row["nome"]), first=self.first)


# ------------------------------------------------------------------------------
# Abas
# ------------------------------------------------------------------------------
class VinculosPanel(ctk.CTkFrame):
    def __init__(self, master, main: "MainFrame"):
        super().__init__(master, fg_color="transparent")
        self.main = main
        self.vinculos: list[dict] = []
        label(self, "Cada ficha 21 comporta 3 vínculos; acima disso, páginas adicionais são geradas "
                    "automaticamente.", wraplength=820, justify="left", anchor="w").pack(anchor="w", pady=(0, 4))
        box = TreeBox(self, ("cbo", "desc", "entrada", "deslig"),
                      (("cbo", "CBO", 80), ("desc", "Descrição", 280), ("entrada", "Entrada", 90),
                       ("deslig", "Desligamento", 100)), height=4)
        box.pack(fill="x", pady=4)
        self.tree = box.tree
        self.tree.bind("<<TreeviewSelect>>", self.on_select)
        self.scroll = ScrollFrame(self, height=330)
        self.scroll.pack(fill="both", expand=True)
        self.form = FormPanel(self.scroll, F21_VINC, per_row=4)
        self.form.pack(anchor="w")
        row = ctk.CTkFrame(self, fg_color="transparent")
        row.pack(anchor="w", pady=6)
        button(row, "Adicionar vínculo", self.add).pack(side="left")
        button(row, "Atualizar selecionado", self.update_sel, kind="secondary").pack(side="left", padx=6)
        button(row, "Remover selecionado", self.remove, kind="danger").pack(side="left")
        button(row, "Limpar campos", self.form.clear, kind="secondary").pack(side="left", padx=6)
        button(row, "Buscar CBO…", self.buscar_cbo, kind="secondary").pack(side="left", padx=(16, 0))

    def refresh(self):
        self.tree.delete(*self.tree.get_children())
        for i, v in enumerate(self.vinculos):
            self.tree.insert("", "end", iid=str(i), values=(
                v.get("cbo_cod", ""), v.get("cbo_desc", ""), v.get("data_entrada", ""), v.get("data_deslig", "")))

    def set_all(self, vinculos: list[dict]):
        self.vinculos = [dict(v) for v in vinculos]
        self.form.clear()
        self.refresh()

    def on_select(self, _=None):
        sel = self.tree.selection()
        if sel:
            self.form.set(self.vinculos[int(sel[0])])

    def _check(self) -> dict | None:
        v = self.form.get()
        if not any(v.values()):
            messagebox.showinfo("Vínculo", "Preencha ao menos um campo do vínculo.", parent=self)
            return None
        return v

    def add(self):
        v = self._check()
        if v:
            self.vinculos.append(v)
            self.refresh()
            self.form.clear()

    def update_sel(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showinfo("Vínculo", "Selecione um vínculo na lista.", parent=self)
            return
        v = self._check()
        if v:
            self.vinculos[int(sel[0])] = v
            self.refresh()

    def remove(self):
        sel = self.tree.selection()
        if sel:
            del self.vinculos[int(sel[0])]
            self.refresh()
            self.form.clear()

    def buscar_cbo(self):
        base = self.main.app.cnes_base
        if not base.available():
            messagebox.showinfo("Base CNES", "Importe a base CNES (aba “Base CNES”) para buscar CBO.", parent=self)
            return
        found = CboSearch(self, base).run()
        if found:
            self.form.vars["cbo_cod"].set(found[0])
            self.form.vars["cbo_desc"].set(found[1])


class RegistroTab(ctk.CTkFrame):
    T_IDENT, T_RESID, T_VINC = "Ficha 20 – Identificação", "Ficha 20 – Residência e banco", "Ficha 21 – Vínculos"

    def __init__(self, master, main: "MainFrame"):
        super().__init__(master, fg_color="transparent")
        self.main = main
        self.editing_id: int | None = None
        today = date.today()
        self.op = tk.StringVar(value="INCLUSAO")
        self.mes = tk.StringVar(value=MESES[today.month - 1].capitalize())
        self.ano = tk.StringVar(value=str(today.year))
        self.envio = tk.StringVar(value=format_date(today))
        self.usa20, self.usa21 = tk.BooleanVar(value=True), tk.BooleanVar(value=False)
        self.descr, self.nome = tk.StringVar(), tk.StringVar()
        _auto_date(self.envio)

        top = ctk.CTkFrame(self)
        top.pack(fill="x")
        label(top, "Dados da operação", font=font(14, True)).pack(anchor="w", padx=12, pady=(8, 2))
        r0 = ctk.CTkFrame(top, fg_color="transparent")
        r0.pack(fill="x", padx=12, pady=2)
        label(r0, "Operação:").pack(side="left")
        for k, t in OPERACOES.items():
            radio(r0, t, k, self.op).pack(side="left", padx=(8, 0))
        label(r0, "Competência:").pack(side="left", padx=(20, 4))
        combo(r0, self.mes, [m.capitalize() for m in MESES], 11, state="readonly").pack(side="left")
        entry(r0, self.ano, 5).pack(side="left", padx=4)
        label(r0, "Envio à regional:").pack(side="left", padx=(16, 4))
        entry(r0, self.envio, 11).pack(side="left")
        r1 = ctk.CTkFrame(top, fg_color="transparent")
        r1.pack(fill="x", padx=12, pady=4)
        label(r1, "Nome do profissional:").pack(side="left")
        entry(r1, self.nome, 44).pack(side="left", padx=6)
        check(r1, "Preencher Ficha 20", self.usa20).pack(side="left", padx=(16, 0))
        check(r1, "Preencher Ficha 21", self.usa21).pack(side="left", padx=(10, 0))
        r2 = ctk.CTkFrame(top, fg_color="transparent")
        r2.pack(fill="x", padx=12, pady=(2, 10))
        label(r2, "Descrição da alteração (vai para a planilha; ex.: Horário Ambulatorial):").pack(side="left")
        entry(r2, self.descr, 34).pack(side="left", padx=6)

        self.nb = ctk.CTkTabview(self)
        self.nb.pack(fill="both", expand=True, pady=8)
        t1, t2, t3 = (self.nb.add(self.T_IDENT), self.nb.add(self.T_RESID), self.nb.add(self.T_VINC))
        sf1 = ScrollFrame(t1)
        sf1.pack(fill="both", expand=True)
        self.f20a = FormPanel(sf1.inner, F20_IDENT, per_row=4)
        self.f20a.pack(anchor="w", padx=4, pady=4)
        button(sf1.inner, "Preencher IBGE/UF de nascimento pelo município (base CNES)",
               lambda: self.lookup_mun("mun_nasc", "ibge_nasc", "uf_nasc"), kind="secondary"
               ).pack(anchor="w", padx=10, pady=8)
        sf2 = ScrollFrame(t2)
        sf2.pack(fill="both", expand=True)
        self.f20b = FormPanel(sf2.inner, F20_RESID + F20_BANCO, per_row=4)
        self.f20b.pack(anchor="w", padx=4, pady=4)
        button(sf2.inner, "Preencher IBGE/UF de residência pelo município (base CNES)",
               lambda: self.lookup_mun("mun_res", "ibge_res", "uf_res"), kind="secondary"
               ).pack(anchor="w", padx=10, pady=8)
        self.vinc = VinculosPanel(t3, main)
        self.vinc.pack(fill="both", expand=True)

        bar = ctk.CTkFrame(self, fg_color="transparent")
        bar.pack(fill="x")
        button(bar, "Salvar e gerar PDFs", lambda: self.save(True)).pack(side="left")
        button(bar, "Salvar sem gerar", lambda: self.save(False), kind="secondary").pack(side="left", padx=8)
        button(bar, "Novo registro (limpar)", self.reset, kind="secondary").pack(side="left")
        self.lbl_edit = label(bar, "")
        self.lbl_edit.pack(side="right")

    # -- helpers
    def lookup_mun(self, nome_key: str, ibge_key: str, uf_key: str):
        base = self.main.app.cnes_base
        panel = self.f20a if nome_key == "mun_nasc" else self.f20b
        if not base.available():
            messagebox.showinfo("Base CNES", "Importe a base CNES (aba “Base CNES”) para usar esta busca.", parent=self)
            return
        m = base.municipio_por_nome(panel.vars[nome_key].get(), panel.vars[uf_key].get().strip())
        if not m:
            messagebox.showinfo("Município", "Município não encontrado. Confira o nome e a UF.", parent=self)
            return
        panel.vars[ibge_key].set(m["codigo"])
        if m.get("uf"):
            panel.vars[uf_key].set(m["uf"])

    def collect(self) -> dict:
        dados = {**self.f20a.get(), **self.f20b.get(), "nome": self.nome.get().strip()}
        try:
            ano = int(self.ano.get())
        except ValueError:
            ano = date.today().year
        mes = [m.capitalize() for m in MESES].index(self.mes.get()) + 1
        return {
            "categoria": "PROFISSIONAL", "operacao": self.op.get(), "ano": ano, "mes": mes,
            "data_envio": self.envio.get().strip(), "usa_ficha20": self.usa20.get(),
            "usa_ficha21": self.usa21.get(), "descricao_planilha": self.descr.get().strip(),
            "dados": dados, "vinculos": self.vinc.vinculos,
        }

    def load(self, form: dict):
        self.editing_id = form["registro_id"]
        self.op.set(form["operacao"])
        self.mes.set(MESES[form["mes"] - 1].capitalize())
        self.ano.set(str(form["ano"]))
        self.envio.set(form["data_envio"])
        self.usa20.set(form["usa_ficha20"])
        self.usa21.set(form["usa_ficha21"])
        self.descr.set(form["descricao_planilha"])
        self.nome.set(form["dados"].get("nome", ""))
        self.f20a.set(form["dados"])
        self.f20b.set(form["dados"])
        self.vinc.set_all(form["vinculos"])
        self.lbl_edit.configure(text=f"Editando registro #{self.editing_id}")

    def reset(self):
        self.editing_id = None
        self.nome.set("")
        self.descr.set("")
        self.f20a.clear()
        self.f20b.clear()
        self.vinc.set_all([])
        self.envio.set(format_date(date.today()))
        self.lbl_edit.configure(text="")

    def save(self, generate: bool):
        est_id = self.main.current_estab_id()
        if est_id is None:
            messagebox.showwarning("Estabelecimento", "Cadastre/selecione o estabelecimento primeiro.", parent=self)
            return
        svc, sess = self.main.app.service, self.main.app.session
        form = self.collect()
        erros, avisos = svc.validar(form)
        if erros:
            messagebox.showerror("Corrija os dados", "\n".join(f"• {e}" for e in erros), parent=self)
            return
        if avisos and not messagebox.askyesno(
                "Atenção", "\n".join(f"• {a}" for a in avisos) + "\n\nSalvar mesmo assim?", parent=self):
            return
        try:
            self.editing_id = svc.salvar(sess, est_id, form, self.editing_id)
            self.lbl_edit.configure(text=f"Editando registro #{self.editing_id}")
            files = svc.gerar_pdfs(sess, self.editing_id) if generate else []
        except Exception as e:  # erro inesperado de E/S, PDF etc.
            messagebox.showerror("Erro", f"Não foi possível concluir:\n{e}", parent=self)
            return
        self.main.refresh_lists()
        if files:
            if messagebox.askyesno("Fichas geradas",
                                   "Arquivos gerados:\n" + "\n".join(str(f) for f in files) +
                                   "\n\nAbrir a pasta?", parent=self):
                open_path(files[0].parent)
        else:
            messagebox.showinfo("Salvo", "Registro salvo.", parent=self)



class HistoricoTab(ctk.CTkFrame):
    def __init__(self, master, main: "MainFrame"):
        super().__init__(master, fg_color="transparent")
        self.main = main
        box = TreeBox(self, ("id", "comp", "op", "nome", "fichas", "envio"),
                      (("id", "Nº", 50), ("comp", "Competência", 100), ("op", "Operação", 90),
                       ("nome", "Profissional / descrição", 360), ("fichas", "Fichas", 90),
                       ("envio", "Envio", 90)), height=14)
        box.pack(fill="both", expand=True)
        self.tree = box.tree
        bar = ctk.CTkFrame(self, fg_color="transparent")
        bar.pack(fill="x", pady=8)
        for text, cmd, kind in (("Abrir para editar", self.edit, "primary"),
                                ("Gerar PDFs novamente", self.regen, "secondary"),
                                ("Abrir pasta de saída", self.open_folder, "secondary"),
                                ("Excluir", self.delete, "danger")):
            button(bar, text, cmd, kind=kind).pack(side="left", padx=(0, 8))

    def refresh(self):
        self.tree.delete(*self.tree.get_children())
        est_id = self.main.current_estab_id()
        if est_id is None:
            return
        for r in self.main.app.db.list_registros(self.main.app.session.user_id, est_id):
            fichas = "+".join(n for n, u in (("20", r["usa_ficha20"]), ("21", r["usa_ficha21"])) if u) or "—"
            if r["categoria"] == "ESTABELECIMENTO":
                nome, fichas = f"[Estabelecimento] {r['descricao_planilha']}", "—"
            else:
                nome = r["nome_profissional"]
            envio = format_date(date.fromisoformat(r["data_envio"])) if r["data_envio"] else ""
            self.tree.insert("", "end", iid=str(r["id"]), values=(
                r["id"], r["competencia"], OPERACOES[r["operacao"]], nome, fichas, envio))

    def _sel(self) -> int | None:
        sel = self.tree.selection()
        if not sel:
            messagebox.showinfo("Histórico", "Selecione um registro.", parent=self)
            return None
        return int(sel[0])

    def edit(self):
        rid = self._sel()
        if rid is None:
            return
        form = self.main.app.service.registro_para_form(self.main.app.session, rid)
        if form["categoria"] == "ESTABELECIMENTO":
            messagebox.showinfo("Histórico", "Este item é só da planilha. Exclua e registre novamente "
                                "para alterar.", parent=self)
            return
        self.main.registro.load(form)
        self.main.tabs.set(MainFrame.TAB_REGISTRO)

    def regen(self):
        rid = self._sel()
        if rid is None:
            return
        try:
            files = self.main.app.service.gerar_pdfs(self.main.app.session, rid)
        except Exception as e:
            messagebox.showerror("Erro", str(e), parent=self)
            return
        if not files:
            messagebox.showinfo("PDFs", "Este registro não tem fichas marcadas.", parent=self)
        elif messagebox.askyesno("Fichas geradas", "\n".join(str(f) for f in files) + "\n\nAbrir a pasta?", parent=self):
            open_path(files[0].parent)

    def open_folder(self):
        est = self.main.current_estab()
        if est:
            p = self.main.app.service.out_root / est["cnes"]
            p.mkdir(parents=True, exist_ok=True)
            open_path(p)

    def delete(self):
        rid = self._sel()
        if rid is not None and messagebox.askyesno("Excluir", "Excluir este registro? Os PDFs já gerados "
                                                   "não são apagados.", parent=self):
            self.main.app.db.delete_registro(self.main.app.session.user_id, rid)
            if self.main.registro.editing_id == rid:
                self.main.registro.reset()
            self.main.refresh_lists()


class PlanilhaTab(ctk.CTkFrame):
    def __init__(self, master, main: "MainFrame"):
        super().__init__(master, fg_color="transparent")
        self.main = main
        today = date.today()
        self.mes = tk.StringVar(value=MESES[today.month - 1].capitalize())
        self.ano = tk.StringVar(value=str(today.year))
        self.obs = tk.StringVar()
        self.anexar = tk.BooleanVar(value=False)
        top = ctk.CTkFrame(self, fg_color="transparent")
        top.pack(fill="x")
        label(top, "Competência:").pack(side="left")
        combo(top, self.mes, [m.capitalize() for m in MESES], 11, state="readonly",
              command=lambda v: self.refresh()).pack(side="left", padx=6)
        ano = entry(top, self.ano, 5)
        ano.pack(side="left")
        ano.bind("<Return>", lambda e: self.refresh())
        ano.bind("<FocusOut>", lambda e: self.refresh())
        button(top, "Atualizar", self.refresh, kind="secondary").pack(side="left", padx=10)
        button(top, "Registrar alteração do estabelecimento…", self.add_estab_change).pack(side="right")
        box = TreeBox(self, ("coluna", "texto", "envio"),
                      (("coluna", "Coluna da planilha", 160), ("texto", "Conteúdo", 460), ("envio", "Envio", 90)),
                      height=14)
        box.pack(fill="both", expand=True, pady=8)
        self.tree = box.tree
        self.lbl = label(self, "", anchor="w")
        self.lbl.pack(anchor="w")
        row = ctk.CTkFrame(self, fg_color="transparent")
        row.pack(fill="x", pady=6)
        label(row, "Observações:").pack(side="left")
        entry(row, self.obs, 64).pack(side="left", padx=6)
        row2 = ctk.CTkFrame(self, fg_color="transparent")
        row2.pack(fill="x")
        check(row2, "Adicionar como nova aba em uma planilha .xlsx existente", self.anexar).pack(side="left")
        button(row2, "Gerar planilha…", self.export).pack(side="right")

    def _comp(self) -> tuple[int, int]:
        try:
            ano = int(self.ano.get())
        except ValueError:
            ano = date.today().year
        return ano, [m.capitalize() for m in MESES].index(self.mes.get()) + 1

    def refresh(self):
        self.tree.delete(*self.tree.get_children())
        est_id = self.main.current_estab_id()
        if est_id is None:
            return
        ano, mes = self._comp()
        linhas = self.main.app.service.linhas_planilha(self.main.app.session, est_id, ano, mes)
        for ln in linhas:
            for key, nome in (("estab", "Estabelecimento"), ("inclusao", "Profissionais – Inclusão"),
                              ("exclusao", "Profissionais – Exclusão"), ("outras", "Outras alterações")):
                if ln.get(key):
                    d = ln.get("data_envio")
                    self.tree.insert("", "end", values=(nome, ln[key], format_date(d) if d else ""))
        self.lbl.configure(text=f"{len(linhas)} linha(s) para {MESES[mes - 1].capitalize()}/{ano}, "
                                f"feitas por {self.main.app.session.nome}.")

    def add_estab_change(self):
        est_id = self.main.current_estab_id()
        if est_id is None:
            return
        win = Modal(self, "Alteração do estabelecimento")
        body = ctk.CTkFrame(win, fg_color="transparent")
        body.pack(padx=16, pady=14)
        txt, env = tk.StringVar(), tk.StringVar(value=format_date(date.today()))
        _auto_date(env)
        label(body, "Descreva a alteração (coluna ESTABELECIMENTO da planilha):").pack(anchor="w")
        entry(body, txt, 56).pack(anchor="w", pady=(4, 8))
        label(body, "Data de envio à regional:").pack(anchor="w")
        entry(body, env, 12).pack(anchor="w", pady=(4, 8))

        def ok():
            form = {"categoria": "ESTABELECIMENTO", "operacao": "ALTERACAO", "ano": self._comp()[0],
                    "mes": self._comp()[1], "data_envio": env.get(), "descricao_planilha": txt.get(),
                    "usa_ficha20": False, "usa_ficha21": False, "dados": {}, "vinculos": []}
            erros, _ = self.main.app.service.validar(form)
            if erros:
                messagebox.showerror("Corrija", "\n".join(erros), parent=win)
                return
            self.main.app.service.salvar(self.main.app.session, est_id, form)
            win.destroy()
            self.main.refresh_lists()
        button(body, "Salvar", ok).pack(anchor="e", pady=(8, 0))
        win.run()

    def export(self):
        est = self.main.current_estab()
        if est is None:
            messagebox.showwarning("Estabelecimento", "Cadastre/selecione o estabelecimento primeiro.", parent=self)
            return
        svc, sess = self.main.app.service, self.main.app.session
        ano, mes = self._comp()
        if not svc.linhas_planilha(sess, est["id"], ano, mes):
            messagebox.showinfo("Planilha", "Não há registros neste mês para este estabelecimento.", parent=self)
            return
        anexar = self.anexar.get()
        if anexar:
            path = filedialog.askopenfilename(parent=self, title="Planilha .xlsx existente",
                                              filetypes=[("Excel", "*.xlsx")])
        else:
            path = filedialog.asksaveasfilename(
                parent=self, title="Salvar planilha", defaultextension=".xlsx",
                initialdir=str(svc.out_root), initialfile=svc.nome_padrao_planilha(est, ano, mes),
                filetypes=[("Excel", "*.xlsx")])
        if not path:
            return
        if anexar and planilha.sheet_name(ano, mes) in _sheetnames(path):
            if not messagebox.askyesno("Aba existente", f"A aba “{planilha.sheet_name(ano, mes)}” já existe "
                                       "nesse arquivo. Substituir?", parent=self):
                return
        try:
            out, replaced = svc.exportar_planilha(sess, est["id"], ano, mes, Path(path), self.obs.get(), anexar)
        except Exception as e:
            messagebox.showerror("Erro", f"Não foi possível gerar a planilha:\n{e}", parent=self)
            return
        extra = ""
        if planilha.libreoffice_disponivel() and messagebox.askyesno(
                "Formato .xls", "Gerar também uma cópia em .xls (formato antigo do Excel)?", parent=self):
            try:
                xls = planilha.converter_para_xls(out)
                extra = f"\n{xls}" if xls else ""
            except Exception as e:
                extra = f"\n(Falha ao converter para .xls: {e})"
        if messagebox.askyesno("Planilha gerada", f"{out}{extra}\n\nAbrir a pasta?", parent=self):
            open_path(out.parent)


def _sheetnames(path: str) -> list[str]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        names = wb.sheetnames
        wb.close()
        return names
    except Exception:
        return []


class BaseTab(ctk.CTkFrame):
    def __init__(self, master, main: "MainFrame"):
        super().__init__(master, fg_color="transparent")
        self.main = main
        self.q: queue.Queue = queue.Queue()
        self.cancel = threading.Event()
        label(self, "Isso é opcional. Importe o ZIP obtido em https://cnes.datasus.gov.br/pages/downloads/arquivosBaseDados.jsp"
                    "(ou https://cnes.datasus.gov.br → Downloads → Base de Dados) para buscar estabelecimento "
                    "pelo CNES, município (IBGE/UF) e CBO por nome. Só as tabelas necessárias são lidas, direto do ZIP. "
                    "Sem a base, tudo funciona com digitação manual.",
                    wraplength=820, justify="left", anchor="w").pack(anchor="w")
        self.info = label(self, "", justify="left", anchor="w")
        self.info.pack(anchor="w", pady=12)
        self.btn = button(self, "Importar ZIP do CNES…", self.start)
        self.btn.pack(anchor="w")
        self.bar = ctk.CTkProgressBar(self, width=420)
        self.bar.set(0)
        self.bar.pack(anchor="w", pady=12)
        self.status = label(self, "", anchor="w")
        self.status.pack(anchor="w")
        self.refresh()

    def refresh(self):
        base = self.main.app.cnes_base
        if base.available():
            i = base.info()
            self.info.configure(text=(
                f"Base instalada: {i.get('importado_de', '?')} (importada em {i.get('importado_em', '?')})\n"
                f"Estabelecimentos: {i.get('n_estabelecimento', 0):,}   Municípios: {i.get('n_municipio', 0):,}   "
                f"CBO: {i.get('n_cbo', 0):,}").replace(",", "."))
        else:
            self.info.configure(text="Nenhuma base CNES importada.")

    def start(self):
        path = filedialog.askopenfilename(parent=self, title="ZIP da base do CNES",
                                          filetypes=[("ZIP", "*.zip")])
        if not path:
            return
        self.btn.configure(state="disabled")
        self.status.configure(text="Importando…")

        def work():
            try:
                rep = cnes_import.import_zip(path, cnes_db_path(),
                                             progress=lambda t, f: self.q.put(("p", t, f)))
                self.q.put(("done", rep))
            except Exception as e:
                self.q.put(("err", str(e)))
        threading.Thread(target=work, daemon=True).start()
        self.after(150, self.poll)

    def poll(self):
        finished = False
        try:
            while True:
                m = self.q.get_nowait()
                if m[0] == "p":
                    self.bar.set(max(0.0, min(1.0, m[2])))
                    self.status.configure(text=f"Importando {m[1]}… {m[2] * 100:.0f}%")
                elif m[0] == "done":
                    finished = True
                    self._done(m[1])
                else:
                    finished = True
                    self.btn.configure(state="normal")
                    self.status.configure(text="")
                    messagebox.showerror("Importação", f"Falha na importação:\n{m[1]}", parent=self)
        except queue.Empty:
            pass
        if not finished:
            self.after(150, self.poll)

    def _done(self, rep: dict):
        self.btn.configure(state="normal")
        self.bar.set(1)
        self.status.configure(text="Concluído.")
        lines = [f"{t}: {n:,} registros".replace(",", ".") for t, n in rep["tabelas"].items()]
        if rep["nao_encontradas"]:
            lines.append("Não encontradas no ZIP: " + ", ".join(rep["nao_encontradas"]))
        for t, d in rep["colunas_ausentes"].items():
            if "faltando" in d:
                lines.append(f"{t}: colunas obrigatórias ausentes {d['faltando']} (cabeçalho: {d['cabecalho'][:8]}…)")
        messagebox.showinfo("Importação", "\n".join(lines), parent=self)
        self.refresh()


# ------------------------------------------------------------------------------
# Janela principal
# ------------------------------------------------------------------------------
class MainFrame(ctk.CTkFrame):
    TAB_REGISTRO, TAB_HIST, TAB_PLAN, TAB_BASE = "Novo registro", "Histórico", "Planilha mensal", "Base CNES"

    def __init__(self, master, app: "App"):
        super().__init__(master, fg_color="transparent")
        self.app = app
        self.estabs: list[dict] = []
        self.est_labels: list[str] = []
        bar = ctk.CTkFrame(self)
        bar.pack(fill="x", padx=8, pady=(8, 4))
        label(bar, f"Usuário: {app.session.nome}", font=font(13, True)).pack(side="left", padx=(12, 0), pady=8)
        label(bar, "Estabelecimento:").pack(side="left", padx=(24, 6))
        self.est_var = tk.StringVar()
        self.est_cb = ctk.CTkComboBox(bar, variable=self.est_var, values=[""], state="readonly", width=430,
                                      command=lambda v: self.refresh_lists())
        self.est_cb.pack(side="left")
        button(bar, "Editar", self.edit_estab, kind="secondary", width=70).pack(side="left", padx=(8, 0))
        button(bar, "Novo", self.new_estab, kind="secondary", width=70).pack(side="left", padx=6)
        button(bar, "Sair", app.logout, kind="danger", width=70).pack(side="right", padx=(8, 12))
        theme_menu(bar, app).pack(side="right")
        label(bar, "Tema:").pack(side="right", padx=(0, 6))

        self.tabs = ctk.CTkTabview(self)
        self.tabs.pack(fill="both", expand=True, padx=8, pady=(4, 8))
        f_reg, f_his, f_pla, f_bas = (self.tabs.add(n) for n in
                                      (self.TAB_REGISTRO, self.TAB_HIST, self.TAB_PLAN, self.TAB_BASE))
        self.registro = RegistroTab(f_reg, self)
        self.historico = HistoricoTab(f_his, self)
        self.planilha = PlanilhaTab(f_pla, self)
        self.base = BaseTab(f_bas, self)
        for tab in (self.registro, self.historico, self.planilha, self.base):
            tab.pack(fill="both", expand=True, padx=4, pady=4)
        self.load_estabs()

    def _est_index(self) -> int:
        v = self.est_var.get()
        return self.est_labels.index(v) if v in self.est_labels else -1

    def load_estabs(self, select_id: int | None = None):
        self.estabs = [dict(r) for r in self.app.db.list_estabelecimentos()]
        self.est_labels = [f"{e['cnes']} – {e['nome_fantasia']}" for e in self.estabs]
        self.est_cb.configure(values=self.est_labels or [""])
        if self.estabs:
            cur = self._est_index()
            idx = next((i for i, e in enumerate(self.estabs) if e["id"] == select_id), 0) \
                if select_id else (cur if cur >= 0 else 0)
            self.est_var.set(self.est_labels[idx])
        else:
            self.est_var.set("")
        self.refresh_lists()

    def current_estab_index(self) -> int:
        return self._est_index()

    def current_estab(self) -> dict | None:
        i = self._est_index()
        return self.estabs[i] if 0 <= i < len(self.estabs) else None

    def current_estab_id(self) -> int | None:
        e = self.current_estab()
        return e["id"] if e else None

    def refresh_lists(self):
        self.historico.refresh()
        self.planilha.refresh()

    def new_estab(self, first: bool = False):
        rid = EstabDialog(self, self, None, first).run()
        if rid:
            self.load_estabs(rid)

    def edit_estab(self):
        e = self.current_estab()
        if not e:
            self.new_estab()
            return
        rid = EstabDialog(self, self, e["id"]).run()
        if rid:
            self.load_estabs(rid)


class App(ctk.CTk):
    def __init__(self, db: Database | None = None):
        self.theme = settings.load().get("theme", "system")
        if self.theme not in THEMES.values():
            self.theme = "system"
        ctk.set_appearance_mode(self.theme)
        ctk.set_default_color_theme("blue")
        super().__init__()
        self.title(f"{APP_TITLE} – {APP_SUBTITLE}  v{__version__}")
        self.geometry("1180x800")
        self.minsize(980, 660)
        set_icon(self)
        style_native(self)
        self.db = db or Database(app_db_path())
        self.cnes_base = cnes_import.CnesBase(cnes_db_path())
        self.service = Service(self.db, self.cnes_base, default_output_dir())
        self.session: Session | None = None
        self.container = ctk.CTkFrame(self, fg_color="transparent")
        self.container.pack(fill="both", expand=True)
        self.show_start()

    def set_theme(self, choice: str) -> None:
        """choice: 'Sistema' | 'Claro' | 'Escuro' (rótulo do menu)."""
        self.theme = THEMES.get(choice, "system")
        ctk.set_appearance_mode(self.theme)
        settings.save(theme=self.theme)
        style_native(self)

    def _clear(self):
        for w in self.container.winfo_children():
            w.destroy()

    def show_start(self):
        self.session = None
        self._clear()
        AuthFrame(self.container, self, first=self.db.user_count() == 0).pack(fill="both", expand=True)

    def on_authenticated(self, session: Session, first: bool = False):
        self.session = session
        self._clear()
        main = MainFrame(self.container, self)
        main.pack(fill="both", expand=True)
        if not main.estabs:
            self.after(200, lambda: main.new_estab(first=True))

    def logout(self):
        if messagebox.askyesno("Sair", "Encerrar a sessão?", parent=self):
            self.show_start()


def run() -> None:
    App().mainloop()
