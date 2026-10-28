"""Teste de fumaça da interface com um tkinter SIMULADO (para ambientes sem display/tkinter).

Executa o fluxo completo pelos handlers reais da GUI. Não substitui um teste visual no Windows.
"""
import sys
import tempfile
import time
import types
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))


# ----------------------------------------------------------------------------- tkinter falso
class Var:
    def __init__(self, master=None, value=""):
        self._v, self._cbs = value, []

    def get(self):
        return self._v

    def set(self, v):
        self._v = v
        for cb in list(self._cbs):
            cb()

    def trace_add(self, mode, cb):
        self._cbs.append(lambda: cb())


class Widget:
    def __init__(self, master=None, *a, **kw):
        self.master, self.kw, self.children, self._items = master, kw, [], {}
        if master is not None and hasattr(master, "children"):
            master.children.append(self)
        self.opts = {}

    def __getattr__(self, name):
        if name.startswith("__"):
            raise AttributeError(name)
        return lambda *a, **k: None

    def winfo_children(self):
        return list(self.children)

    def winfo_toplevel(self):
        return self

    def destroy(self):
        if self.master is not None and self in getattr(self.master, "children", []):
            self.master.children.remove(self)

    def __setitem__(self, k, v):
        self.opts[k] = v

    def __getitem__(self, k):
        return self.opts.get(k)

    def configure(self, **kw):
        self.opts.update(kw)

    def after(self, ms, fn=None, *a):
        self.opts.setdefault("_after", []).append(fn)

    def mainloop(self):
        pass


class Combobox(Widget):
    _cur = -1

    def current(self, i=None):
        if i is None:
            return self._cur
        self._cur = i


class Tabview(Widget):
    def __init__(self, *a, **k):
        super().__init__(*a, **k)
        self.tabs, self.current = {}, None

    def add(self, name):
        self.tabs[name] = Widget(self)
        return self.tabs[name]

    def tab(self, name):
        return self.tabs[name]

    def set(self, name):
        self.current = name


class Tree(Widget):
    def __init__(self, *a, **k):
        super().__init__(*a, **k)
        self.rows, self._sel, self._n = {}, (), 0

    def insert(self, parent, index, iid=None, values=(), **k):
        iid = iid or f"I{self._n}"
        self._n += 1
        self.rows[iid] = values
        return iid

    def delete(self, *ids):
        for i in ids:
            self.rows.pop(i, None)

    def get_children(self, *a):
        return tuple(self.rows)

    def selection(self):
        return self._sel

    def selection_set(self, *ids):
        self._sel = tuple(ids)


def install_fake_tk():
    tk = types.ModuleType("tkinter")
    ttk = types.ModuleType("tkinter.ttk")
    mb = types.ModuleType("tkinter.messagebox")
    fd = types.ModuleType("tkinter.filedialog")
    for n in ("Tk", "Toplevel", "Canvas", "Listbox"):
        setattr(tk, n, type(n, (Widget,), {}))
    tk.StringVar = Var
    tk.BooleanVar = type("BooleanVar", (Var,), {"__init__": lambda s, master=None, value=False: Var.__init__(s, master, value)})
    for n in ("Frame", "Label", "Entry", "Button", "Radiobutton", "Checkbutton", "LabelFrame",
              "Notebook", "Scrollbar", "Spinbox", "Progressbar"):
        setattr(ttk, n, type(n, (Widget,), {}))
    ttk.Combobox, ttk.Treeview = Combobox, Tree
    tk.PhotoImage = Widget
    tk.TclError = Exception
    class _Style:
        def __init__(self, *a, **k): pass
        def theme_use(self, *a, **k): pass
        def configure(self, *a, **k): pass
        def map(self, *a, **k): pass
    ttk.Style = _Style
    # --- customtkinter falso
    ctk = types.ModuleType("customtkinter")
    for n in ("CTk", "CTkToplevel", "CTkFrame", "CTkLabel", "CTkButton", "CTkEntry", "CTkComboBox",
              "CTkCheckBox", "CTkRadioButton", "CTkScrollableFrame", "CTkProgressBar", "CTkOptionMenu",
              "CTkScrollbar", "CTkFont"):
        setattr(ctk, n, type(n, (Widget,), {}))
    ctk.CTkTabview = Tabview
    ctk.MODE = ["system"]
    ctk.set_appearance_mode = lambda m: ctk.MODE.__setitem__(0, m)
    ctk.get_appearance_mode = lambda: "Dark" if ctk.MODE[0] == "dark" else "Light"
    ctk.set_default_color_theme = lambda n: None
    log = {"errors": [], "infos": [], "asks": [], "warnings": []}
    mb.showerror = lambda t, m, **k: log["errors"].append((t, m))
    mb.showinfo = lambda t, m, **k: log["infos"].append((t, m))
    mb.showwarning = lambda t, m, **k: log["warnings"].append((t, m))
    mb.askyesno = lambda t, m, **k: (log["asks"].append((t, m)), False)[1]  # nunca abre pastas
    fd.state = {}
    fd.askopenfilename = lambda **k: fd.state.get("open", "")
    fd.asksaveasfilename = lambda **k: fd.state.get("save", "")
    tk.ttk, tk.messagebox, tk.filedialog = ttk, mb, fd
    sys.modules.update({"customtkinter": ctk, "tkinter": tk, "tkinter.ttk": ttk, "tkinter.messagebox": mb,
                        "tkinter.filedialog": fd})
    return log, fd, ctk


LOG, FD, CTK = install_fake_tk()

from app import gui  # noqa: E402  (depois do tkinter falso)
from app.db import Database  # noqa: E402

failures = []


def check(cond, msg):
    print(("  ok  " if cond else " FAIL ") + msg)
    if not cond:
        failures.append(msg)


def main():
    tmp = Path(tempfile.mkdtemp())
    db = Database(tmp / "app.db")
    gui.cnes_db_path = lambda: tmp / "cnes.db"
    app = gui.App(db)
    app.cnes_base.path = tmp / "cnes.db"
    app.service.out_root = tmp / "saida"

    print("1) primeiro acesso")
    auth = app.container.children[0]
    check(isinstance(auth, gui.AuthFrame) and auth.first, "tela de criação do primeiro usuário")
    auth.nome.set("Maria Juliani"); auth.user.set("maria"); auth.pw.set("curta"); auth.pw2.set("curta")
    auth.submit()
    check(len(LOG["errors"]) == 1 and db.user_count() == 0, "senha curta é recusada")
    auth.pw.set("senha-forte-1"); auth.pw2.set("diferente")
    auth.submit()
    check(db.user_count() == 0, "confirmação diferente é recusada")
    auth.pw2.set("senha-forte-1")
    auth.submit()
    check(app.session and app.session.nome == "Maria Juliani", "usuário criado e logado")
    main_f = app.container.children[0]
    check(isinstance(main_f, gui.MainFrame) and main_f.estabs == [], "tela principal sem estabelecimento")

    print("2) estabelecimento")
    reg = main_f.registro
    reg.nome.set("X")
    reg.save(False)
    check(any("estabelecimento" in w[1].lower() for w in LOG["warnings"]), "bloqueia registro sem estabelecimento")
    dlg = gui.EstabDialog(main_f, main_f, None, True)
    dlg.salvar()
    check("CNES" in LOG["errors"][-1][1], "CNES obrigatório (7 dígitos)")
    dlg.form.set({"cnes": "2733579", "nome_fantasia": ""})
    dlg.salvar()
    check("nome fantasia" in LOG["errors"][-1][1].lower(), "nome fantasia obrigatório")
    dlg.form.set({"cnes": "2733579", "nome_fantasia": "HOSPITAL REGIONAL CRISTO REI",
                  "municipio": "Astorga", "uf": "PR", "ibge": "4102000"})
    dlg.salvar()
    main_f.load_estabs(dlg.result)
    check(main_f.current_estab()["cnes"] == "2733579", "estabelecimento selecionado")

    print("3) registro com Ficha 20 + 21")
    reg.nome.set("Maria Aparecida da Silva Souza")
    reg.op.set("INCLUSAO")
    reg.usa20.set(True); reg.usa21.set(True)
    reg.f20a.set({"cpf": "52998224725", "data_nasc": "05031985", "sexo": "F", "cadastramento": "SUS"})
    check(reg.f20a.vars["data_nasc"].get() == "05/03/1985", "data digitada sem barras é formatada")
    reg.f20a.vars["raca_cor"].set("01 - Branca")
    check(reg.f20a.get()["raca_cor"] == "01", "combo devolve só o código")
    reg.f20b.set({"logradouro": "Rua das Flores", "cep": "86730000"})
    reg.vinc.form.set({"cbo_cod": "223505", "cbo_desc": "Enfermeiro", "data_entrada": "01/03/2024", "sus": "S"})
    reg.vinc.add()
    check(len(reg.vinc.vinculos) == 1 and reg.vinc.tree.rows, "vínculo adicionado à lista")
    reg.f20a.vars["cpf"].set("52998224724")
    n_ask = len(LOG["asks"])
    reg.save(True)
    check(len(LOG["asks"]) == n_ask + 1 and not list((tmp / "saida").rglob("*.pdf")),
          "CPF inválido gera aviso e (sem confirmar) não salva")
    reg.f20a.vars["cpf"].set("52998224725")
    reg.save(True)
    pdfs = sorted((tmp / "saida" / "2733579").rglob("*.pdf"))
    check(len(pdfs) == 2, f"gerou Ficha 20 e Ficha 21 ({[p.name for p in pdfs]})")
    check(reg.editing_id is not None, "registro salvo")

    print("4) só Ficha 21 / só Ficha 20")
    reg.reset()
    reg.nome.set("Joao Teste"); reg.op.set("EXCLUSAO")
    reg.usa20.set(False); reg.usa21.set(True)
    reg.vinc.form.set({"cbo_cod": "225125", "data_deslig": "31/10/2026"}); reg.vinc.add()
    reg.save(True)
    check(len(list((tmp / "saida").rglob("Ficha21_Joao*"))) == 1 and
          not list((tmp / "saida").rglob("Ficha20_Joao*")), "apenas Ficha 21 gerada")
    reg.reset()
    reg.nome.set("Ana Alterada"); reg.op.set("ALTERACAO"); reg.descr.set("Horário Ambulatorial")
    reg.usa20.set(True); reg.usa21.set(False)
    reg.save(True)
    check(len(list((tmp / "saida").rglob("Ficha20_Ana*"))) == 1, "apenas Ficha 20 gerada")
    reg.reset()
    reg.usa20.set(False); reg.usa21.set(False); reg.nome.set("Sem Ficha")
    n_err = len(LOG["errors"])
    reg.save(True)
    check(len(LOG["errors"]) == n_err + 1, "exige ao menos uma ficha")

    print("5) histórico e edição")
    main_f.refresh_lists()
    hist = main_f.historico
    check(len(hist.tree.rows) == 3, f"histórico lista 3 registros ({len(hist.tree.rows)})")
    first = list(hist.tree.rows)[-1]
    hist.tree.selection_set(first)
    hist.edit()
    check(reg.nome.get() == "Maria Aparecida da Silva Souza" and reg.f20a.vars["cpf"].get() == "52998224725"
          and len(reg.vinc.vinculos) == 1, "registro carregado para edição com vínculos")
    hist.regen()

    print("6) planilha")
    plan = main_f.planilha
    plan.add_estab_change  # existe
    form = {"categoria": "ESTABELECIMENTO", "operacao": "ALTERACAO", "ano": plan._comp()[0], "mes": plan._comp()[1],
            "data_envio": "27/10/2026", "descricao_planilha": "Novo telefone", "usa_ficha20": False,
            "usa_ficha21": False, "dados": {}, "vinculos": []}
    app.service.salvar(app.session, main_f.current_estab_id(), form)
    plan.refresh()
    cols = [v[0] for v in plan.tree.rows.values()]
    check("Profissionais – Inclusão" in cols and "Profissionais – Exclusão" in cols
          and "Outras alterações" in cols and "Estabelecimento" in cols, f"preview com as 4 colunas {sorted(set(cols))}")
    FD.state["save"] = str(tmp / "saida" / "plan.xlsx")
    plan.export()
    import openpyxl
    ws = openpyxl.load_workbook(tmp / "saida" / "plan.xlsx").active
    vals = [c.value for row in ws.iter_rows(min_row=7, max_row=12) for c in row if c.value]
    check("Maria Aparecida da Silva Souza" in vals and "Joao Teste" in vals, "xlsx contém inclusão e exclusão")
    check(any(isinstance(v, str) and v.startswith("Ana Alterada (Horário") for v in vals), "xlsx contém alteração com descrição")
    check(ws["A3"].value == "MUNICÍPIO: ASTORGA" and "MARIA JULIANI" in str(ws["A27"].value), "município e responsável")

    print("7) base CNES (importação em thread)")
    z = tmp / "cnes.zip"
    sys.path.insert(0, str(Path(__file__).parent))
    from test_core import _make_zip
    _make_zip(z)
    FD.state["open"] = str(z)
    base_tab = main_f.base
    base_tab.start()
    t0 = time.time()
    while time.time() - t0 < 20:
        base_tab.poll()
        if base_tab.status.opts.get("text") == "Concluído.":
            break
        time.sleep(0.1)
    check(app.cnes_base.available(), "base importada e disponível")
    dlg = gui.EstabDialog(main_f, main_f, None, False)
    dlg.form.vars["cnes"].set("19")
    dlg.buscar()
    check(dlg.form.vars["nome_fantasia"].get() == "UNIDADE BÁSICA JOSÉ" and dlg.form.vars["municipio"].get() == "ASTORGA",
          "busca de estabelecimento pelo CNES preenche o formulário")
    reg.f20b.vars["mun_res"].set("astorga"); reg.f20b.vars["uf_res"].set("")
    reg.lookup_mun("mun_res", "ibge_res", "uf_res")
    check(reg.f20b.vars["ibge_res"].get() == "4102000" and reg.f20b.vars["uf_res"].get() == "PR", "IBGE/UF por município")

    print("7b) tema e abas")
    app.set_theme("Escuro")
    check(app.theme == "dark" and CTK.get_appearance_mode() == "Dark", "tema escuro aplicado")
    check(gui.tree_colors()["bg"] == "#2b2b2b", "cores da tabela seguem o tema escuro")
    app.set_theme("Claro")
    check(app.theme == "light" and gui.tree_colors()["bg"] == "#ffffff", "tema claro aplicado")
    from app import settings
    check(settings.load().get("theme") == "light", "tema salvo em config.json")
    app.set_theme("Sistema")
    hist.tree.selection_set(list(hist.tree.rows)[-1]); hist.edit()
    check(main_f.tabs.current == gui.MainFrame.TAB_REGISTRO, "editar do histórico volta à aba de registro")

    print("8) logout e login")
    app.show_start()
    auth = app.container.children[0]
    check(isinstance(auth, gui.AuthFrame) and not auth.first, "tela de login (não é mais primeiro acesso)")
    auth.user.set("maria"); auth.pw.set("errada"); n_err = len(LOG["errors"]); auth.submit()
    check(len(LOG["errors"]) == n_err + 1 and app.session is None, "senha errada recusada")
    auth.pw.set("senha-forte-1"); auth.submit()
    check(app.session is not None and app.container.children[0].estabs, "login correto reabre com estabelecimento salvo")

    print("\nFALHAS:", len(failures))
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
