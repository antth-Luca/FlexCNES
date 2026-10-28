import io
import os
import sys
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from pypdf import PdfReader

from app import cnes_import, pdf_fill, planilha, validators
from app.db import AuthError, Database
from app.security import hash_password, verify_password

ESTAB = {"cnes": "2733579", "nome_fantasia": "HOSPITAL REGIONAL CRISTO REI"}


class SecurityTests(unittest.TestCase):
    def test_hash_roundtrip(self):
        h = hash_password("segredo123", iterations=1000)
        self.assertTrue(verify_password("segredo123", h))
        self.assertFalse(verify_password("outra", h))
        self.assertNotEqual(h, hash_password("segredo123", iterations=1000))  # salt aleatório
        self.assertFalse(verify_password("x", "lixo"))


class ValidatorTests(unittest.TestCase):
    def test_cpf(self):
        self.assertTrue(validators.valid_cpf("529.982.247-25"))
        self.assertFalse(validators.valid_cpf("111.111.111-11"))
        self.assertFalse(validators.valid_cpf("529.982.247-24"))

    def test_cns(self):
        self.assertTrue(validators.valid_cns("898001160660001"))   # provisório
        self.assertTrue(validators.valid_cns("700000000000005"))
        self.assertFalse(validators.valid_cns("898001160660002"))

    def test_pis(self):
        self.assertTrue(validators.valid_pis("12345678919"))
        self.assertFalse(validators.valid_pis("12345678901"))

    def test_dates(self):
        self.assertEqual(validators.parse_date("05/03/1985"), date(1985, 3, 5))
        self.assertEqual(validators.parse_date("05031985"), date(1985, 3, 5))
        self.assertIsNone(validators.parse_date("31/02/2020"))
        self.assertEqual(validators.date_digits("5/3/1985"), "05031985")
        self.assertEqual(validators.date_digits("05/03/1985"), "05031985")


class DbTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = Database(Path(self.tmp.name) / "t.db")

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def test_first_access_and_login(self):
        self.assertEqual(self.db.user_count(), 0)
        uid = self.db.create_user("maria", "Maria Juliani", "senha-forte-1")
        self.assertEqual(self.db.user_count(), 1)
        self.assertEqual(self.db.authenticate("MARIA", "senha-forte-1")["id"], uid)  # login sem caixa
        with self.assertRaises(AuthError):
            self.db.authenticate("maria", "errada")
        with self.assertRaises(AuthError):
            self.db.authenticate("fulano", "senha-forte-1")
        with self.assertRaises(AuthError):
            self.db.create_user("Maria", "Outra", "senha-forte-2")  # duplicado

    def test_estabelecimento_create_and_edit(self):
        eid = self.db.save_estabelecimento({"cnes": "2733579", "nome_fantasia": "Hospital X",
                                            "municipio": "Astorga", "uf": "pr"})
        self.db.save_estabelecimento({"cnes": "2733579", "nome_fantasia": "Hospital Y",
                                      "municipio": "Astorga", "uf": "PR"}, est_id=eid)
        e = self.db.get_estabelecimento(eid)
        self.assertEqual((e["nome_fantasia"], e["uf"]), ("Hospital Y", "PR"))
        with self.assertRaises(ValueError):
            self.db.save_estabelecimento({"cnes": "123", "nome_fantasia": "x"})

    def test_registros_isolados_por_usuario(self):
        u1 = self.db.create_user("a", "A", "senha-forte-1")
        u2 = self.db.create_user("b", "B", "senha-forte-1")
        eid = self.db.save_estabelecimento({"cnes": "2733579", "nome_fantasia": "H"})
        rid = self.db.save_registro(u1, eid, operacao="INCLUSAO", competencia="2026-10",
                                    data_envio="2026-10-27", nome_profissional="Fulana",
                                    dados={"nome": "Fulana"})
        self.assertEqual(len(self.db.list_registros(u1)), 1)
        self.assertEqual(len(self.db.list_registros(u2)), 0)
        self.assertIsNone(self.db.get_registro(u2, rid))
        with self.assertRaises(AuthError):
            self.db.save_registro(u2, eid, operacao="EXCLUSAO", competencia="2026-10",
                                  data_envio=None, registro_id=rid)
        self.db.delete_registro(u2, rid)  # não apaga o de outro usuário
        self.assertEqual(len(self.db.list_registros(u1)), 1)
        self.db.delete_registro(u1, rid)
        self.assertEqual(len(self.db.list_registros(u1)), 0)


class PdfTests(unittest.TestCase):
    def test_ficha20_texto_e_paginas(self):
        d = {"nome": "Maria Souza", "cpf": "529.982.247-25", "data_nasc": "05/03/1985",
             "sexo": "F", "cadastramento": "SUS", "logradouro": "Rua das Flores",
             "livro": "A12"}
        with tempfile.TemporaryDirectory() as t:
            out = pdf_fill.write_ficha20(Path(t) / "f20.pdf", ESTAB, "INCLUSAO", d)
            r = PdfReader(str(out))
            self.assertEqual(len(r.pages), 1)
            txt = r.pages[0].extract_text()
        self.assertIn("MARIA SOUZA", txt)
        self.assertIn("RUA DAS FLORES", txt)
        self.assertIn("HOSPITAL REGIONAL CRISTO REI", txt)
        self.assertIn("FICHA CADASTRAL", txt)  # conteúdo original preservado
        self.assertIn("A", txt)

    def test_ficha21_paginas_por_vinculos(self):
        v = {"cbo_cod": "223505", "cbo_desc": "Enfermeiro", "sus": "S"}
        with tempfile.TemporaryDirectory() as t:
            for n, esperado in ((0, 1), (1, 1), (3, 1), (4, 2), (7, 3)):
                out = pdf_fill.write_ficha21(Path(t) / f"f21_{n}.pdf", ESTAB, "ALTERACAO",
                                             "Maria", [v] * n)
                self.assertEqual(len(PdfReader(str(out)).pages), esperado, n)

    def test_vinculos_incluem_conteudo(self):
        v = {"cbo_cod": "223505", "cbo_desc": "Enfermeiro", "data_entrada": "01/03/2024"}
        with tempfile.TemporaryDirectory() as t:
            out = pdf_fill.write_ficha21(Path(t) / "f.pdf", ESTAB, "INCLUSAO", "Maria", [v])
            txt = PdfReader(str(out)).pages[0].extract_text()
        self.assertIn("ENFERMEIRO", txt)
        self.assertIn("MARIA", txt)


class PlanilhaTests(unittest.TestCase):
    def _regs(self, n_inc=3):
        d = {"cpf": "04854630914", "cns": "703205645229093", "categoria_prof": "Auxiliar de Enfermagem"}
        regs = [dict(categoria="PROFISSIONAL", operacao="INCLUSAO", nome_profissional=f"Pessoa {i}",
                     data_envio="2026-10-27", descricao_planilha="", dados=d) for i in range(n_inc)]
        regs.append(dict(categoria="PROFISSIONAL", operacao="EXCLUSAO", nome_profissional="Saiu",
                         data_envio="2026-10-27", descricao_planilha="",
                         dados={"cpf": "11228170967", "cns": "709806041946495",
                                "vinculos": [{"cbo_desc": "MEDICO CLINICO", "cbo_cod": "225125"}]}))
        regs.append(dict(categoria="PROFISSIONAL", operacao="ALTERACAO", nome_profissional="Mudou",
                         data_envio="2026-10-27", descricao_planilha="Horário Ambulatorial", dados={}))
        regs.append(dict(categoria="ESTABELECIMENTO", operacao="ALTERACAO", nome_profissional="",
                         data_envio="2026-10-27", descricao_planilha="Novo telefone"))
        return regs

    def test_colunas_e_rodape(self):
        linhas = planilha.montar_linhas(self._regs(), ESTAB, "OUTUBRO")
        with tempfile.TemporaryDirectory() as t:
            p, rep = planilha.write_workbook(Path(t) / "p.xlsx", ano=2026, mes=10,
                                             municipio="Astorga", uf="pr", regional="15ª Regional de Saúde", linhas=linhas,
                                             realizado_por="Maria", observacoes="obs")
            wb = openpyxl.load_workbook(p)
        self.assertEqual(wb.sheetnames, ["Outubro2026"])
        ws = wb.active
        self.assertEqual(ws["A2"].value, "MUNICÍPIO: ASTORGA - PR          15ª Regional de Saúde")
        self.assertEqual(ws["A6"].value, 2733579)
        self.assertEqual(ws["E6"].value, "Novo telefone")
        self.assertEqual((ws["F7"].value, ws["G7"].value), ("Pessoa 0", "Auxiliar de Enfermagem"))
        self.assertEqual((ws["H7"].value, ws["I7"].value), ("04854630914", "703205645229093"))
        self.assertEqual(ws["J10"].value, "Saiu")
        self.assertEqual((ws["K10"].value, ws["L10"].value, ws["M10"].value),
                         ("Medico Clinico", "11228170967", "709806041946495"))
        self.assertEqual(ws["N11"].value, "Mudou – Horário Ambulatorial")
        self.assertEqual(ws["D6"].value.date(), date(2026, 10, 27))
        # dados 6..24 (19 linhas mínimas); rodapé começa na linha 25
        self.assertEqual(ws["A25"].value, "ALTERAÇÕES FEITA POR: MARIA")
        self.assertEqual(ws["A26"].value, "OBSERVAÇÕES:  obs")
        self.assertIn("A25:N25", [str(r) for r in ws.merged_cells.ranges])
        self.assertEqual(ws.print_area, "'Outubro2026'!$A$1:$N$28")

    def test_titulo_sem_regional_ou_uf(self):
        self.assertEqual(planilha._titulo_municipio("Astorga"), "MUNICÍPIO: ASTORGA")
        self.assertEqual(planilha._titulo_municipio("", "", "X Regional"), "MUNICÍPIO:" + " " * 10 + "X Regional")

    def test_alteracao_com_cpf_cns_no_texto(self):
        regs = [dict(categoria="PROFISSIONAL", operacao="ALTERACAO", nome_profissional="Ana",
                     data_envio="2026-10-27", descricao_planilha="CBO X para Y",
                     dados={"cpf": "123", "cns": "456"})]
        ln = planilha.montar_linhas(regs, ESTAB, "OUTUBRO")[0]
        self.assertEqual(ln["outras"], "Ana (CPF 123; CNS 456) – CBO X para Y")

    def test_rodape_desce_com_muitas_linhas(self):
        linhas = planilha.montar_linhas(self._regs(n_inc=30), ESTAB, "OUTUBRO")  # 33 linhas
        with tempfile.TemporaryDirectory() as t:
            p, _ = planilha.write_workbook(Path(t) / "p.xlsx", ano=2026, mes=10, municipio="Astorga",
                                           linhas=linhas, realizado_por="Maria")
            ws = openpyxl.load_workbook(p).active
        self.assertEqual(ws["A24"].value, 2733579)   # ainda é linha de dados
        self.assertEqual(ws["A" + str(25 + len(linhas) - 19)].value, "ALTERAÇÕES FEITA POR: MARIA")

    def test_anexar_e_substituir(self):
        linhas = planilha.montar_linhas(self._regs(), ESTAB, "OUTUBRO")
        with tempfile.TemporaryDirectory() as t:
            f = Path(t) / "p.xlsx"
            planilha.write_workbook(f, ano=2026, mes=9, municipio="X", linhas=linhas, realizado_por="M")
            _, r1 = planilha.write_workbook(f, ano=2026, mes=10, municipio="X", linhas=linhas,
                                            realizado_por="M", anexar_existente=True)
            _, r2 = planilha.write_workbook(f, ano=2026, mes=10, municipio="X", linhas=linhas[:2],
                                            realizado_por="M", anexar_existente=True)
            self.assertEqual((r1, r2), (False, True))
            self.assertEqual(openpyxl.load_workbook(f).sheetnames, ["Outubro2026", "Setembro2026"])


def _make_zip(path, encoding="cp1252"):
    est = ("CO_UNIDADE;CO_CNES;NO_RAZAO_SOCIAL;NO_FANTASIA;CO_MUNICIPIO_GESTOR;NU_CNPJ;"
           "NO_LOGRADOURO;NU_ENDERECO;NO_BAIRRO;CO_CEP;NU_TELEFONE\n"
           "1;2733579;ASSOC. CRISTO REI;HOSPITAL REGIONAL CRISTO REI;410200;12345678000199;"
           "RUA A;10;CENTRO;86730000;4432341234\n"
           "2;0000019;SECRETARIA;UNIDADE BÁSICA JOSÉ;410200;;;;;;\n")
    mun = "CO_MUNICIPIO;NO_MUNICIPIO;CO_SIGLA_ESTADO\n4102000;ASTORGA;PR\n4106902;CURITIBA;PR\n"
    cbo = ("CO_CBO;DS_ATIVIDADE_PROFISSIONAL\n223505;ENFERMEIRO\n225125;MÉDICO CLÍNICO\n"
           "322205;TÉCNICO DE ENFERMAGEM\n")
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("tbEstabelecimento202610.csv", est.encode(encoding))
        z.writestr("tbMunicipio202610.csv", mun.encode(encoding))
        z.writestr("tbAtividadeProfissional202610.csv", cbo.encode(encoding))
        z.writestr("rlEstabEquipe202610.csv", "x;y\n1;2\n")  # tabela ignorada


class CnesTests(unittest.TestCase):
    def _run(self, encoding):
        with tempfile.TemporaryDirectory() as t:
            z, db = Path(t) / "c.zip", Path(t) / "cnes.db"
            _make_zip(z, encoding)
            seen = []
            rep = cnes_import.import_zip(z, db, progress=lambda tb, f: seen.append((tb, f)))
            base = cnes_import.CnesBase(db)
            self.assertTrue(base.available())
            self.assertEqual(rep["tabelas"], {"estabelecimento": 2, "municipio": 2, "cbo": 3})
            self.assertEqual(rep["nao_encontradas"], [])
            e = base.estabelecimento("2733579")
            self.assertEqual(e["nome_fantasia" if "nome_fantasia" in e else "fantasia"],
                             "HOSPITAL REGIONAL CRISTO REI")
            self.assertEqual((e["municipio"], e["uf"], e["ibge"]), ("ASTORGA", "PR", "4102000"))
            self.assertEqual(base.estabelecimento("19")["fantasia"], "UNIDADE BÁSICA JOSÉ")  # zfill
            self.assertIsNone(base.estabelecimento("9999999"))
            self.assertEqual(base.cbo("223505"), "ENFERMEIRO")
            self.assertEqual(base.search_cbo("medico")[0][0], "225125")   # ignora acento
            self.assertEqual(base.municipio_por_nome("Astorga", "PR")["nome"], "ASTORGA")
            self.assertEqual(base.municipio_por_nome("astorga", "PR")["codigo"], "4102000")
            self.assertGreater(len(seen), 0)

    def test_importa_latin1(self):
        self._run("cp1252")

    def test_importa_utf8(self):
        self._run("utf-8")

    def test_zip_sem_tabelas_esperadas(self):
        with tempfile.TemporaryDirectory() as t:
            z = Path(t) / "c.zip"
            with zipfile.ZipFile(z, "w") as zz:
                zz.writestr("outra.csv", "a;b\n1;2\n")
            rep = cnes_import.import_zip(z, Path(t) / "c.db")
            self.assertEqual(sorted(rep["nao_encontradas"]), ["cbo", "estabelecimento", "municipio"])
            self.assertFalse(cnes_import.CnesBase(Path(t) / "c.db").available())


class RegionalDbTests(unittest.TestCase):
    def test_regional_salva_e_migra_banco_antigo(self):
        import sqlite3
        with tempfile.TemporaryDirectory() as t:
            f = Path(t) / "velho.db"
            c = sqlite3.connect(f)   # banco antigo, sem a coluna "regional"
            c.executescript("""CREATE TABLE estabelecimentos (id INTEGER PRIMARY KEY AUTOINCREMENT,
                cnes TEXT NOT NULL UNIQUE, nome_fantasia TEXT NOT NULL, razao_social TEXT DEFAULT '',
                cnpj TEXT DEFAULT '', municipio TEXT DEFAULT '', uf TEXT DEFAULT '', ibge TEXT DEFAULT '',
                atualizado_em TEXT NOT NULL);""")
            c.commit(); c.close()
            db = Database(f)
            eid = db.save_estabelecimento({**ESTAB, "regional": " 15ª Regional de Saúde "})
            self.assertEqual(db.get_estabelecimento(eid)["regional"], "15ª Regional de Saúde")
            db.close()


class IbgeValidationTests(unittest.TestCase):
    def _erros(self, ibge):
        from app.service import Service
        svc = Service.__new__(Service)
        form = {"usa_ficha20": True, "dados": {"ibge_nasc": ibge, "ibge_res": ibge}}
        erros, _ = Service.validar(svc, form)
        return [e for e in erros if "IBGE" in e]

    def test_ibge_6_ou_7_digitos(self):
        self.assertEqual(self._erros("410190"), [])      # padrão CNES (6)
        self.assertEqual(self._erros("4101900"), [])     # com dígito verificador (7)
        self.assertEqual(self._erros(""), [])
        self.assertEqual(len(self._erros("41019")), 2)   # 5 dígitos: inválido
        self.assertEqual(len(self._erros("41019000")), 2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
